"""Arety Office — a local, offline, WPS-style office suite for the browser.

Runs on 127.0.0.1 only. Documents are edited in the browser (Quill / x-spreadsheet);
files on disk are read and written through pandoc, LibreOffice and openpyxl.
"""
import os, re, io, csv, json, time, shutil, hashlib, subprocess, tempfile, pathlib, urllib.parse
from fastapi import FastAPI, HTTPException, Request
from fastapi.responses import FileResponse, JSONResponse, RedirectResponse, HTMLResponse
from fastapi.staticfiles import StaticFiles
from pydantic import BaseModel

HOME = pathlib.Path.home()
BASE = pathlib.Path(__file__).resolve().parent
STATIC = BASE / "static"
CONF = HOME / ".config/officesuite"; CONF.mkdir(parents=True, exist_ok=True)
CACHE = HOME / ".cache/officesuite"; CACHE.mkdir(parents=True, exist_ok=True)
MEDIA = CACHE / "media"; MEDIA.mkdir(exist_ok=True)
RECENT = CONF / "recent.json"
ALLOWED_ROOTS = [HOME, pathlib.Path("/media"), pathlib.Path("/mnt"), pathlib.Path("/tmp")]
CONVERT_DOC = HOME / "bin/convert-doc"
WPS_OFFLINE = HOME / "bin/wps-offline"

DOC = {"docx", "doc", "odt", "rtf", "md", "markdown", "txt", "html", "htm", "wps", "epub"}
SHEET = {"xlsx", "xls", "csv", "ods", "et"}
SLIDES = {"pptx", "ppt", "odp", "dps"}
PDF = {"pdf"}
OFFICE = DOC | SHEET | SLIDES | PDF

app = FastAPI(title="Arety Office")

# ---------- helpers ----------
def ext(p): return pathlib.Path(p).suffix.lstrip(".").lower()
def kind_of(p):
    e = ext(p)
    return "doc" if e in DOC else "sheet" if e in SHEET else "slides" if e in SLIDES else "pdf" if e in PDF else None

def safe(p: str) -> pathlib.Path:
    path = pathlib.Path(os.path.expanduser(p)).resolve()
    if not any(path == r or r in path.parents for r in ALLOWED_ROOTS):
        raise HTTPException(403, f"{path} is outside the allowed folders")
    return path

def run(*cmd, timeout=180):
    r = subprocess.run(cmd, stdout=subprocess.PIPE, stderr=subprocess.STDOUT, text=True, timeout=timeout,
                       env={**os.environ, "HOME": str(HOME)})
    if r.returncode: raise HTTPException(500, f"{cmd[0]} failed:\n{r.stdout[-1500:]}")
    return r.stdout

def soffice_convert(src: pathlib.Path, fmt: str, outdir: pathlib.Path, infilter=None) -> pathlib.Path:
    args = ["soffice", "--headless", "--norestore"]
    if infilter: args.append(f"--infilter={infilter}")
    run(*args, "--convert-to", fmt, "--outdir", str(outdir), str(src))
    out = outdir / (src.stem + "." + fmt.split(":")[0])
    if not out.exists(): raise HTTPException(500, f"LibreOffice produced no .{fmt} for {src.name}")
    return out

def load_recent():
    try: return json.loads(RECENT.read_text())
    except Exception: return []

def touch_recent(path: pathlib.Path):
    items = [r for r in load_recent() if r["path"] != str(path)]
    items.insert(0, {"path": str(path), "name": path.name, "kind": kind_of(path), "ts": time.time()})
    RECENT.write_text(json.dumps(items[:60], indent=1))

def file_info(p: pathlib.Path):
    st = p.stat()
    return {"name": p.name, "path": str(p), "dir": p.is_dir(), "kind": None if p.is_dir() else kind_of(p),
            "size": 0 if p.is_dir() else st.st_size, "mtime": st.st_mtime}

# ---------- static / pages ----------
@app.get("/", response_class=HTMLResponse)
def index(): return (STATIC / "index.html").read_text()

@app.get("/open")
def open_redirect(path: str):                 # used by the desktop file handler
    return RedirectResponse("/#open=" + urllib.parse.quote(str(safe(path))))

@app.get("/file")
def get_file(path: str):
    p = safe(path)
    if not p.is_file(): raise HTTPException(404, "no such file")
    return FileResponse(str(p), filename=p.name, content_disposition_type="inline")

# ---------- home page data ----------
@app.get("/api/recent")
def api_recent():
    items = [r for r in load_recent() if pathlib.Path(r["path"]).exists()]
    return items

@app.post("/api/recent/remove")
def api_recent_remove(body: dict):
    RECENT.write_text(json.dumps([r for r in load_recent() if r["path"] != body.get("path")], indent=1)); return {"ok": True}

@app.get("/api/browse")
def api_browse(path: str = str(HOME), all: bool = False):
    p = safe(path)
    if not p.is_dir(): raise HTTPException(404, "not a folder")
    items = []
    for c in sorted(p.iterdir(), key=lambda x: (not x.is_dir(), x.name.lower())):
        if c.name.startswith("."): continue
        try:
            if c.is_dir() or all or ext(c) in OFFICE: items.append(file_info(c))
        except OSError: pass
    return {"path": str(p), "parent": str(p.parent) if p != p.parent else None, "items": items,
            "places": [{"name": n, "path": str(HOME / d)} for n, d in
                       [("Home", ""), ("Documents", "Documents"), ("Desktop", "Desktop"), ("Downloads", "Downloads")]
                       if (HOME / d).is_dir()] + [{"name": "Drives", "path": "/media"}] * pathlib.Path("/media").is_dir()}

@app.get("/api/search")
def api_search(q: str, path: str = str(HOME)):
    root = safe(path); q = q.lower(); out = []
    for dp, dns, fns in os.walk(root):
        dns[:] = [d for d in dns if not d.startswith(".") and d not in ("node_modules", "snap", "AppData")]
        for f in fns:
            if q in f.lower() and ext(f) in OFFICE:
                out.append(file_info(pathlib.Path(dp) / f))
                if len(out) >= 200: return out
    return out

# ---------- create ----------
class NewFile(BaseModel):
    kind: str; name: str; dir: str = str(HOME / "Documents")

@app.post("/api/new")
def api_new(body: NewFile):
    d = safe(body.dir); d.mkdir(parents=True, exist_ok=True)
    name = body.name.strip() or "Untitled"
    default_ext = {"doc": "docx", "sheet": "xlsx", "slides": "pptx"}[body.kind]
    if ext(name) not in OFFICE: name += "." + default_ext
    target = d / name; n = 1
    while target.exists():
        n += 1; target = d / f"{pathlib.Path(name).stem} ({n}){pathlib.Path(name).suffix}"
    if body.kind == "doc":
        if ext(target) in ("docx", "odt", "md", "html", "epub"):
            run("pandoc", "-f", "markdown", "-t", ext(target), "-o", str(target), "/dev/null")
        elif ext(target) == "txt": target.write_text("")
        else:
            tmp = target.with_suffix(".docx"); run("pandoc", "-f", "markdown", "-o", str(tmp), "/dev/null")
            soffice_convert(tmp, ext(target), d); tmp.unlink(missing_ok=True)
    elif body.kind == "sheet":
        from openpyxl import Workbook
        if ext(target) == "csv": target.write_text("")
        else:
            wb = Workbook(); tmp = target if ext(target) == "xlsx" else target.with_suffix(".xlsx"); wb.save(tmp)
            if tmp != target: soffice_convert(tmp, ext(target), d); tmp.unlink(missing_ok=True)
    else:
        tmp = target if ext(target) == "pptx" else target.with_suffix(".pptx")
        md = tempfile.NamedTemporaryFile("w", suffix=".md", delete=False); md.write(f"% {target.stem}\n"); md.close()
        run("pandoc", md.name, "-o", str(tmp)); os.unlink(md.name)
        if tmp != target: soffice_convert(tmp, ext(target), d); tmp.unlink(missing_ok=True)
    touch_recent(target); return file_info(target)

# ---------- open ----------
def media_dir_for(path: pathlib.Path):
    h = hashlib.sha1(str(path).encode()).hexdigest()[:16]
    d = MEDIA / h; shutil.rmtree(d, ignore_errors=True); d.mkdir(parents=True); return h, d

def doc_to_html(path: pathlib.Path):
    e = ext(path); src = path; tmpdir = None
    if e in ("doc", "wps", "rtf"):
        tmpdir = pathlib.Path(tempfile.mkdtemp()); src = soffice_convert(path, "docx", tmpdir); e = "docx"
    if e == "txt": return "<p>" + "</p><p>".join(html_escape(l) for l in path.read_text(errors="replace").splitlines()) + "</p>"
    fmt = {"md": "markdown", "markdown": "markdown", "htm": "html"}.get(e, e)
    h, mdir = media_dir_for(path)
    html = run("pandoc", "-f", fmt, "-t", "html", "--extract-media", str(mdir), str(src))
    html = html.replace(f'src="{mdir}/', f'src="/media/{h}/')
    if tmpdir: shutil.rmtree(tmpdir, ignore_errors=True)
    return html

def html_escape(s): return s.replace("&", "&amp;").replace("<", "&lt;").replace(">", "&gt;")

def sheet_to_json(path: pathlib.Path):
    from openpyxl import load_workbook
    from openpyxl.utils import get_column_letter
    e = ext(path); src = path; tmpdir = None
    if e == "csv":
        rows = list(csv.reader(open(path, newline="", encoding="utf-8", errors="replace")))
        return [{"name": path.stem, "rows": {str(r): {"cells": {str(c): {"text": v} for c, v in enumerate(row) if v != ""}}
                                             for r, row in enumerate(rows)}}]
    if e in ("xls", "ods", "et"):
        tmpdir = pathlib.Path(tempfile.mkdtemp()); src = soffice_convert(path, "xlsx", tmpdir)
    wb = load_workbook(src)  # formulas kept as text
    sheets = []
    for ws in wb.worksheets:
        styles, style_idx = [], {}
        rows, merges = {}, []
        for mr in ws.merged_cells.ranges:
            merges.append((mr.min_row - 1, mr.min_col - 1, mr.max_row - mr.min_row, mr.max_col - mr.min_col))
        merge_map = {(r, c): (rs, cs) for r, c, rs, cs in merges}
        for row in ws.iter_rows():
            for cell in row:
                if cell.value is None and not cell.has_style: continue
                v = cell.value
                if v is None: text = ""
                elif isinstance(v, bool): text = "TRUE" if v else "FALSE"
                elif hasattr(v, "isoformat"): text = v.isoformat(sep=" ")
                else: text = str(v)
                st = {}
                f = cell.font
                if f:
                    if f.bold: st.setdefault("font", {})["bold"] = True
                    if f.italic: st.setdefault("font", {})["italic"] = True
                    if f.underline: st["underline"] = True
                    if f.color is not None and f.color.type == "rgb" and isinstance(f.color.rgb, str) and f.color.rgb not in ("FF000000",):
                        st["color"] = "#" + f.color.rgb[-6:]
                fill = cell.fill
                if fill is not None and fill.fill_type == "solid" and fill.fgColor is not None and fill.fgColor.type == "rgb" and isinstance(fill.fgColor.rgb, str) and fill.fgColor.rgb not in ("00000000",):
                    st["bgcolor"] = "#" + fill.fgColor.rgb[-6:]
                if cell.alignment is not None and cell.alignment.horizontal in ("left", "center", "right"):
                    st["align"] = cell.alignment.horizontal
                if text == "" and not st: continue
                c = {"text": text}
                if st:
                    key = json.dumps(st, sort_keys=True)
                    if key not in style_idx: style_idx[key] = len(styles); styles.append(st)
                    c["style"] = style_idx[key]
                rc = (cell.row - 1, cell.column - 1)
                if rc in merge_map: c["merge"] = list(merge_map[rc])
                rows.setdefault(str(rc[0]), {"cells": {}})["cells"][str(rc[1])] = c
        cols = {}
        for i in range(1, (ws.max_column or 1) + 1):
            w = ws.column_dimensions[get_column_letter(i)].width
            if w: cols[str(i - 1)] = {"width": int(w * 7.5)}
        sheets.append({"name": ws.title, "rows": rows, "cols": cols, "styles": styles,
                       "merges": [f"{get_column_letter(c+1)}{r+1}:{get_column_letter(c+cs+1)}{r+rs+1}" for r, c, rs, cs in merges]})
    if tmpdir: shutil.rmtree(tmpdir, ignore_errors=True)
    return sheets

def slides_pdf(path: pathlib.Path):
    h = hashlib.sha1(str(path).encode()).hexdigest()[:16]
    out = CACHE / f"{h}.pdf"
    if not out.exists() or out.stat().st_mtime < path.stat().st_mtime:
        tmpdir = pathlib.Path(tempfile.mkdtemp()); pdf = soffice_convert(path, "pdf", tmpdir)
        shutil.move(str(pdf), out); shutil.rmtree(tmpdir, ignore_errors=True)
    return out

@app.get("/api/open")
def api_open(path: str):
    p = safe(path)
    if not p.is_file(): raise HTTPException(404, "no such file")
    k = kind_of(p)
    if not k: raise HTTPException(415, "not an office file")
    touch_recent(p)
    base = {"path": str(p), "name": p.name, "kind": k, "ext": ext(p)}
    if k == "doc": return {**base, "html": doc_to_html(p)}
    if k == "sheet": return {**base, "sheets": sheet_to_json(p)}
    if k == "slides": return {**base, "pdf": "/file?path=" + urllib.parse.quote(str(slides_pdf(p)))}
    return {**base, "pdf": "/file?path=" + urllib.parse.quote(str(p))}

# ---------- save ----------
class SaveDoc(BaseModel):
    path: str; html: str
class SaveSheet(BaseModel):
    path: str; sheets: list

@app.post("/api/save/doc")
def api_save_doc(body: SaveDoc):
    p = safe(body.path); e = ext(p)
    html = re.sub(r'src="/media/([^"]+)"', lambda m: f'src="{MEDIA}/{m.group(1)}"', body.html)
    tmp = pathlib.Path(tempfile.mkdtemp()); src = tmp / "doc.html"; src.write_text(html)
    if e == "txt":
        p.write_text(run("pandoc", "-f", "html", "-t", "plain", "--wrap=none", str(src)))
    elif e in ("docx", "odt", "md", "markdown", "html", "htm", "epub"):
        fmt = {"md": "gfm", "markdown": "gfm", "htm": "html"}.get(e, e)
        run("pandoc", "-f", "html", "-t", fmt, "-o", str(p), str(src))
    else:  # doc / rtf / wps: go through docx
        mid = tmp / "doc.docx"; run("pandoc", "-f", "html", "-t", "docx", "-o", str(mid), str(src))
        out = soffice_convert(mid, e, tmp); shutil.move(str(out), p)
    shutil.rmtree(tmp, ignore_errors=True); touch_recent(p); return file_info(p)

@app.post("/api/save/sheet")
def api_save_sheet(body: SaveSheet):
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill, Alignment
    from openpyxl.utils import get_column_letter
    p = safe(body.path); e = ext(p)
    if e == "csv":
        s = body.sheets[0] if body.sheets else {"rows": {}}
        rows = s.get("rows", {}); maxr = max([int(r) for r in rows if r.isdigit()] + [-1])
        with open(p, "w", newline="") as f:
            w = csv.writer(f)
            for r in range(maxr + 1):
                cells = rows.get(str(r), {}).get("cells", {}); maxc = max([int(c) for c in cells] + [-1])
                w.writerow([cells.get(str(c), {}).get("text", "") for c in range(maxc + 1)])
        touch_recent(p); return file_info(p)
    wb = Workbook(); wb.remove(wb.active)
    for s in body.sheets:
        ws = wb.create_sheet((s.get("name") or "Sheet")[:31]); styles = s.get("styles", [])
        for r, row in s.get("rows", {}).items():
            if not r.isdigit(): continue
            for c, cell in row.get("cells", {}).items():
                text = cell.get("text", "")
                if text == "" and "style" not in cell: continue
                oc = ws.cell(row=int(r) + 1, column=int(c) + 1)
                if isinstance(text, str) and text.startswith("="): oc.value = text
                else:
                    try: oc.value = int(text) if re.fullmatch(r"-?\d+", str(text)) else float(text) if re.fullmatch(r"-?\d*\.\d+(e-?\d+)?", str(text)) else text
                    except Exception: oc.value = text
                st = styles[cell["style"]] if isinstance(cell.get("style"), int) and cell["style"] < len(styles) else None
                if st:
                    fnt = st.get("font", {})
                    oc.font = Font(bold=fnt.get("bold", False), italic=fnt.get("italic", False), underline="single" if st.get("underline") else None,
                                   color=st["color"].lstrip("#").rjust(8, "F") if st.get("color") else None,
                                   size=fnt.get("size") or None, name=fnt.get("name") or None)
                    if st.get("bgcolor"): oc.fill = PatternFill("solid", fgColor=st["bgcolor"].lstrip("#").rjust(8, "F"))
                    if st.get("align"): oc.alignment = Alignment(horizontal=st["align"])
                if cell.get("merge"):
                    rs, cs = cell["merge"]; ws.merge_cells(start_row=int(r) + 1, start_column=int(c) + 1, end_row=int(r) + 1 + rs, end_column=int(c) + 1 + cs)
        for c, col in (s.get("cols") or {}).items():
            if c.isdigit() and col.get("width"): ws.column_dimensions[get_column_letter(int(c) + 1)].width = col["width"] / 7.5
    if e == "xlsx": wb.save(p)
    else:
        tmp = pathlib.Path(tempfile.mkdtemp()); mid = tmp / (p.stem + ".xlsx"); wb.save(mid)
        out = soffice_convert(mid, e, tmp); shutil.move(str(out), p); shutil.rmtree(tmp, ignore_errors=True)
    touch_recent(p); return file_info(p)

# ---------- tools ----------
@app.post("/api/convert")
def api_convert(body: dict):
    p = safe(body["path"]); fmt = re.sub(r"[^a-z0-9-]", "", body["fmt"].lower())
    flags = ["--ocr"] if body.get("ocr") else []
    out = run(str(CONVERT_DOC), *flags, str(p), fmt, timeout=600).strip().splitlines()[-1]
    return file_info(pathlib.Path(out))

@app.post("/api/export-pdf")
def api_export_pdf(body: dict):
    p = safe(body["path"]); tmp = pathlib.Path(tempfile.mkdtemp()); pdf = soffice_convert(p, "pdf", tmp)
    target = p.with_suffix(".pdf"); shutil.move(str(pdf), target); shutil.rmtree(tmp, ignore_errors=True)
    return file_info(target)

@app.post("/api/open-external")
def api_open_external(body: dict):
    p = safe(body["path"]); k = kind_of(p)
    if body.get("app") == "wps" and WPS_OFFLINE.exists():
        flag = {"doc": "--writer", "sheet": "--sheets", "slides": "--slides", "pdf": "--pdf"}.get(k, "")
        subprocess.Popen([str(WPS_OFFLINE)] + ([flag] if flag else []) + [str(p)], start_new_session=True)
    elif body.get("app") == "folder":
        subprocess.Popen(["xdg-open", str(p.parent)], start_new_session=True)
    else:
        subprocess.Popen(["xdg-open", str(p)], start_new_session=True)
    return {"ok": True}

@app.post("/api/fs")
def api_fs(body: dict):
    op = body.get("op"); p = safe(body["path"])
    if op == "rename":
        new = p.with_name(body["name"]); p.rename(new); return file_info(new)
    if op == "delete":
        if p.is_dir(): shutil.rmtree(p)
        else: p.unlink()
        RECENT.write_text(json.dumps([r for r in load_recent() if r["path"] != str(p)], indent=1)); return {"ok": True}
    if op == "mkdir":
        new = p / body["name"]; new.mkdir(); return file_info(new)
    if op == "copy":
        new = safe(body["to"]); shutil.copy2(p, new); return file_info(new)
    raise HTTPException(400, "unknown op")

@app.post("/api/clientlog")
async def api_clientlog(req: Request):
    import sys; print("[client]", (await req.body()).decode()[:500], file=sys.stderr, flush=True); return {"ok": True}

app.mount("/media", StaticFiles(directory=str(MEDIA)), name="media")
app.mount("/static", StaticFiles(directory=str(STATIC)), name="static")

if __name__ == "__main__":
    import uvicorn
    uvicorn.run(app, host="127.0.0.1", port=int(os.environ.get("OFFICE_PORT", "8792")), log_level="warning")
